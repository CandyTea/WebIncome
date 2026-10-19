"""Validate and export the supplied Sheet2 template without modifying Excel.

This is a local research prototype. It does not push or upload any files.
Only data in A:Y is exported. Red driver names exclude the whole row.
"""

from __future__ import annotations

import argparse
import colorsys
import hashlib
import json
import os
import re
from collections import Counter
from datetime import date, datetime
from decimal import Decimal, ROUND_HALF_UP
from io import BytesIO
from pathlib import Path
from tempfile import NamedTemporaryFile
from xml.etree import ElementTree as ET
from zipfile import ZipFile, BadZipFile

import openpyxl
from openpyxl.styles.colors import COLOR_INDEX
from openpyxl.utils.datetime import from_excel


class ExportError(ValueError):
    pass


DEFAULT_YEAR = 2026


def sheet_date(value, epoch, coordinate):
    """Read the saved Sheet2 header; yearless labels use the query year."""
    try:
        if numeric(value):
            value = from_excel(value, epoch)
        if isinstance(value, datetime):
            return value.date().isoformat()
        if isinstance(value, date):
            return value.isoformat()
        if isinstance(value, str):
            text = value.strip()
            match = re.fullmatch(r"(?:(\d{4})年)?(\d{1,2})月(\d{1,2})日(?:总流水)?", text)
            if match:
                year, month, day = match.groups()
                return date(int(year) if year else DEFAULT_YEAR, int(month), int(day)).isoformat()
            match = re.fullmatch(r"(\d{4})[-/](\d{1,2})[-/](\d{1,2})", text)
            if match:
                return date(*map(int, match.groups())).isoformat()
    except (ValueError, OverflowError):
        pass
    raise ExportError(f"Sheet2!{coordinate} 日期无效，请使用日期或“10月19日总流水”这样的表头。")


def numeric(value):
    return isinstance(value, (int, float, Decimal)) and not isinstance(value, bool)


def money_cents(value):
    amount = Decimal(str(value))
    if not amount.is_finite():
        raise ExportError("流水数值不是有限数。")
    return int((amount * 100).quantize(Decimal("1"), rounding=ROUND_HALF_UP))


def theme_palette(workbook):
    if not workbook.loaded_theme:
        return {}
    ns = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
    scheme = ET.fromstring(workbook.loaded_theme).find("a:themeElements/a:clrScheme", ns)
    if scheme is None:
        return {}
    # Spreadsheet theme indices use light1, dark1, light2, dark2 first.
    names = ["lt1", "dk1", "lt2", "dk2"] + [f"accent{i}" for i in range(1, 7)] + ["hlink", "folHlink"]
    palette = {}
    for index, name in enumerate(names):
        entry = scheme.find(f"a:{name}", ns)
        if entry is not None and len(entry):
            color = entry[0]
            palette[index] = color.get("lastClr") or color.get("val")
    return palette


def font_rgb(color, palette):
    if color is None or color.type == "auto":
        return "000000"
    if color.type == "rgb":
        rgb = color.rgb[-6:]
    elif color.type == "theme":
        rgb = palette.get(color.theme)
    elif color.type == "indexed":
        if color.indexed == 64:
            return "000000"
        if not 0 <= color.indexed < len(COLOR_INDEX):
            raise ExportError("姓名使用了无法识别的索引颜色。")
        rgb = COLOR_INDEX[color.indexed][-6:]
    else:
        rgb = None
    if not rgb or len(rgb) != 6:
        raise ExportError("无法解析姓名字体颜色，请人工检查。")
    channels = [int(rgb[i:i + 2], 16) / 255 for i in (0, 2, 4)]
    if color.tint:
        hue, lightness, saturation = colorsys.rgb_to_hls(*channels)
        tint = color.tint
        lightness = lightness * (1 + tint) if tint < 0 else lightness * (1 - tint) + tint
        channels = colorsys.hls_to_rgb(hue, lightness, saturation)
    return "".join(f"{round(channel * 255):02X}" for channel in channels)


def is_red(rgb):
    channels = [int(rgb[i:i + 2], 16) / 255 for i in (0, 2, 4)]
    hue, saturation, value = colorsys.rgb_to_hsv(*channels)
    # Include red theme shades. This is an explicit prototype rule, not an
    # arbitrary test that any non-black name is red.
    return (hue <= 20 / 360 or hue >= 340 / 360) and saturation >= 0.35 and value >= 0.15


def reject_conditional_name_fonts(sheet):
    for area in sheet.conditional_formatting:
        touches_names = any(r.min_col <= 1 <= r.max_col and r.max_row >= 2 for r in area.sqref.ranges)
        if touches_names and any(rule.dxf is not None and rule.dxf.font is not None for rule in sheet.conditional_formatting[area]):
            raise ExportError("姓名出现条件格式字体规则，需要先核验实际显示颜色。")


def extract(path):
    return extract_bytes(Path(path).read_bytes())


def extract_bytes(source_bytes):
    # Both readers see the identical saved file, even if Excel saves again later.
    try:
        with ZipFile(BytesIO(source_bytes)) as archive:
            if len(archive.infolist()) > 2000 or sum(item.file_size for item in archive.infolist()) > 100 * 1024 * 1024:
                raise ExportError("Excel 解压后过大，请使用正常的流水工作簿。")
    except BadZipFile as error:
        raise ExportError("不是有效的 .xlsx 文件。") from error
    formulas = openpyxl.load_workbook(BytesIO(source_bytes), data_only=False, rich_text=True)
    values = openpyxl.load_workbook(BytesIO(source_bytes), data_only=True)
    try:
        if "Sheet2" not in formulas.sheetnames:
            raise ExportError("需要名为 Sheet2 的工作表。")
        sheet, cached = formulas["Sheet2"], values["Sheet2"]
        if sheet.max_row > 5000:
            raise ExportError("行数超过当前导入器限制，请检查工作表的有效范围。")
        if cached["A1"].value != "姓名":
            raise ExportError("Sheet2!A1 应为姓名表头；请先在 Excel/WPS 中计算并保存。")
        reject_conditional_name_fonts(sheet)
        palette = theme_palette(formulas)
        dates = [sheet_date(cached.cell(1, col).value, values.epoch, cached.cell(1, col).coordinate)
                 for col in range(2, 26)]
        if len(set(dates)) != 24:
            raise ExportError("日期重复，不能生成唯一的每日记录。")

        drivers, records = [], []
        counts, excluded_colors = Counter(), Counter()
        seen_names, excluded_names = set(), set()
        for row in range(2, sheet.max_row + 1):
            name = cached.cell(row, 1).value
            if name is None or name == 0 or name == "":
                if any(cached.cell(row, col).value is not None for col in range(2, 26)):
                    raise ExportError(f"第 {row} 行缺少有效姓名，但有流水数据。")
                counts["invalid_name_rows"] += 1
                continue
            if not isinstance(name, str) or name != name.strip():
                raise ExportError(f"第 {row} 行姓名格式需要人工核对。")
            if name in seen_names:
                raise ExportError(f"第 {row} 行存在重复姓名，不能可靠合并司机。")
            seen_names.add(name)
            if not isinstance(sheet.cell(row, 1).value, (str, int, float, type(None))):
                raise ExportError(f"第 {row} 行姓名有富文本或特殊公式，需要核验颜色。")
            rgb = font_rgb(sheet.cell(row, 1).font.color, palette)
            if is_red(rgb):
                counts["excluded_red_drivers"] += 1
                excluded_colors[rgb] += 1
                excluded_names.add(name)
                continue
            driver_id = "driver-" + hashlib.sha256(name.encode("utf-8")).hexdigest()[:16]
            drivers.append({"id": driver_id, "name": name})
            counts["included_drivers"] += 1
            if sheet.row_dimensions[row].hidden:
                counts["included_hidden_drivers"] += 1
            for col, day in enumerate(dates, 2):
                cell, cache = sheet.cell(row, col), cached.cell(row, col)
                amount = cache.value
                if cache.data_type == "e":
                    raise ExportError(f"Sheet2!{cell.coordinate} 包含计算错误。")
                if cell.data_type == "f":
                    if not numeric(amount):
                        raise ExportError(f"Sheet2!{cell.coordinate} 缺少数值缓存，请计算并保存 Excel。")
                    counts["cached_formula_cells"] += 1
                elif cell.value is not None and not numeric(amount):
                    raise ExportError(f"Sheet2!{cell.coordinate} 含非数值流水。")
                if amount is None:
                    cents = None
                    counts["blank_cells"] += 1
                elif numeric(amount):
                    cents = money_cents(amount)
                    counts["numeric_cells"] += 1
                    counts["zero_cells" if cents == 0 else "nonzero_cells"] += 1
                    counts["total_amount_cents"] += cents
                else:
                    raise ExportError(f"Sheet2!{cell.coordinate} 不是数值或空白。")
                records.append({"driverId": driver_id, "date": day, "amountCents": cents})

        if not drivers:
            raise ExportError("没有可导出的司机，保留旧快照供人工核查。")
        if excluded_names.intersection(driver["name"] for driver in drivers):
            raise ExportError("导出包含被排除的姓名。")
        if len({(record["driverId"], record["date"]) for record in records}) != len(records):
            raise ExportError("司机与日期记录重复。")
        payload = {
            "schemaVersion": 1,
            "sourceRange": "Sheet2!A:Y",
            "amountUnit": "CNY-cent",
            "dates": dates,
            "drivers": drivers,
            "records": records,
        }
        # Audit contains counts only; excluded identities are never serialized.
        audit = {
            "sourceSha256": hashlib.sha256(source_bytes).hexdigest(),
            "dateFrom": min(dates),
            "dateTo": max(dates),
            "dateColumns": len(dates),
            **dict(counts),
            "excluded_font_colors": dict(excluded_colors),
            "record_count": len(records),
        }
        return payload, audit
    finally:
        formulas.close()
        values.close()


def atomic_json(path, payload):
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    temp_path = None
    try:
        with NamedTemporaryFile("w", encoding="utf-8", dir=path.parent, suffix=".tmp", delete=False) as handle:
            temp_path = Path(handle.name)
            json.dump(payload, handle, ensure_ascii=False, indent=2, allow_nan=False)
            handle.write("\n")
        os.replace(temp_path, path)
    finally:
        if temp_path is not None and temp_path.exists():
            temp_path.unlink()


def main():
    parser = argparse.ArgumentParser(description="只读取 Sheet2!A:Y 的日期与已保存数值，整行排除红色姓名。")
    parser.add_argument("workbook", type=Path)
    parser.add_argument("--output", type=Path, default=Path(".local/sheet2.json"))
    parser.add_argument("--audit", type=Path, default=Path(".local/validation.json"))
    args = parser.parse_args()
    paths = [args.workbook.resolve(), args.output.resolve(), args.audit.resolve()]
    if len(set(paths)) != len(paths) or any(p.suffix.lower() in (".xlsx", ".xls") for p in paths[1:]):
        parser.error("输出文件必须与源文件不同，且不能是 Excel 文件。")
    try:
        payload, audit = extract(args.workbook)
        atomic_json(args.output, payload)
        atomic_json(args.audit, audit)
    except (ExportError, OSError, ValueError) as error:
        parser.exit(1, f"导出失败：{error}\n")
    print(json.dumps(audit, ensure_ascii=False, indent=2))
    print(f"本地数据：{args.output.resolve()}")


if __name__ == "__main__":
    main()
